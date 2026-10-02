import json
import subprocess
from pathlib import Path

import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent.parent.parent

DATASET_PATH = (
    BASE_DIR
    / "synthetic_data"
    / "synthetic_dataset.xlsx"
)

if __import__("sys").platform.startswith("win"):
    OPA_PATH = BASE_DIR / "opa" / "opa.exe"
else:
    OPA_PATH = BASE_DIR / "opa"
REGO_PATH = (
    BASE_DIR
    / "policies"
    / "rego"
    / "policy.rego"
)

RESULTS_DIR = (
    BASE_DIR
    / "evaluation_results"
)

RESULTS_PATH = (
    RESULTS_DIR
    / "results.json"
)

RESULTS_EXCEL_PATH = (
    RESULTS_DIR
    / "results.xlsx"
)

EXPECTED_COLUMNS = [
    "record_id",
    "customer_name",
    "email",
    "phone",
    "address",
    "dob",
    "gender",
    "passport_number",
    "ni_number",
    "credit_card_number",
    "bank_account",
    "medical_condition",
    "ethnicity",
    "religion",
    "political_view",
    "employee_id",
    "department",
    "job_role",
    "customer_id",
    "ip_address",
    "feedback",
]


def validate_paths():
    """
    Make sure all required files exist.
    """

    required_paths = {
        "Dataset": DATASET_PATH,
        "OPA": OPA_PATH,
        "Rego policy": REGO_PATH,
    }

    for name, path in required_paths.items():

        if not path.exists():
            raise FileNotFoundError(
                f"{name} not found: {path}"
            )


def load_dataset():
    """
    Load the synthetic Excel dataset.
    """

    print()
    print("Loading synthetic dataset...")
    print(f"Dataset: {DATASET_PATH}")

    dataframe = pd.read_excel(
        DATASET_PATH
    )

    print(
        f"Rows    : {len(dataframe)}"
    )

    print(
        f"Columns : {len(dataframe.columns)}"
    )

    return dataframe


def validate_dataset_columns(dataframe):
    """
    Verify that the dataset contains exactly the
    expected project columns.
    """

    actual_columns = list(
        dataframe.columns
    )

    missing_columns = [
        column
        for column in EXPECTED_COLUMNS
        if column not in actual_columns
    ]

    unexpected_columns = [
        column
        for column in actual_columns
        if column not in EXPECTED_COLUMNS
    ]

    if missing_columns:
        raise ValueError(
            "Dataset is missing required columns: "
            + ", ".join(missing_columns)
        )

    if unexpected_columns:
        raise ValueError(
            "Dataset contains unexpected columns: "
            + ", ".join(unexpected_columns)
        )

    print(
        "Dataset column validation: PASSED"
    )


def clean_value(value):
    """
    Convert pandas/Excel values into JSON-safe values.

    Empty Excel cells become empty strings.
    Numeric values that represent whole numbers are converted
    without the unnecessary .0 suffix.
    """

    if pd.isna(value):
        return ""

    if isinstance(value, pd.Timestamp):
        return value.isoformat()

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value)


def row_to_input(row):
    """
    Convert one pandas row into the JSON object that
    will be passed to OPA.
    """

    record = {}

    for column in EXPECTED_COLUMNS:

        record[column] = clean_value(
            row[column]
        )

    return record


def evaluate_record(record):
    """
    Evaluate one record using OPA.
    """

    input_json = json.dumps(
        record,
        ensure_ascii=False,
    )

    command = [
        str(OPA_PATH),
        "eval",
        "-d",
        str(REGO_PATH),
        "--stdin-input",
        "--format",
        "json",
        "data.policy.result",
    ]

    process = subprocess.run(
        command,
        input=input_json,
        text=True,
        capture_output=True,
        encoding="utf-8",
    )

    if process.returncode != 0:

        raise RuntimeError(
            "OPA evaluation failed:\n"
            + process.stderr
        )

    output = json.loads(
        process.stdout
    )

    try:
        result = (
            output["result"][0]
            ["expressions"][0]
            ["value"]
        )
    except (
        KeyError,
        IndexError,
        TypeError,
    ) as error:

        raise RuntimeError(
            "Unexpected OPA response:\n"
            + json.dumps(
                output,
                indent=2,
            )
        ) from error

    return result


def build_record_result(record, evaluation):
    """
    Combine the original record information with
    the OPA evaluation result.
    """

    triggered_rules = evaluation.get(
        "triggered_rules",
        [],
    )

    rule_ids = [
        rule["rule_id"]
        for rule in triggered_rules
    ]

    categories = [
        rule["category"]
        for rule in triggered_rules
    ]

    return {
        "record_id": record["record_id"],
        "outcome": evaluation["outcome"],
        "triggered_rule_ids": rule_ids,
        "triggered_categories": categories,
        "triggered_rules": triggered_rules,
        "input": record,
    }


def evaluate_dataset(dataframe):
    """
    Evaluate every dataset record through OPA.
    """

    results = []

    total_records = len(dataframe)

    print()
    print("=" * 60)
    print("STARTING DATASET EVALUATION")
    print("=" * 60)

    for index, (_, row) in enumerate(
        dataframe.iterrows(),
        start=1,
    ):

        record = row_to_input(row)

        evaluation = evaluate_record(
            record
        )

        result = build_record_result(
            record,
            evaluation,
        )

        results.append(result)

        print(
            f"[{index}/{total_records}] "
            f"Record {record['record_id']} "
            f"→ {evaluation['outcome']}"
        )

    return results


def calculate_summary(results):
    """
    Calculate basic PASS / FLAG / BLOCK statistics.
    """

    total = len(results)

    pass_count = sum(
        1
        for result in results
        if result["outcome"] == "PASS"
    )

    flag_count = sum(
        1
        for result in results
        if result["outcome"] == "FLAG"
    )

    block_count = sum(
        1
        for result in results
        if result["outcome"] == "BLOCK"
    )

    pass_rate = (
        (pass_count / total) * 100
        if total
        else 0
    )

    return {
        "total_records": total,
        "pass": pass_count,
        "flag": flag_count,
        "block": block_count,
        "pass_rate": round(
            pass_rate,
            2,
        ),
    }


def save_results(results, summary):
    """
    Save the complete evaluation results.
    """

    RESULTS_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    output = {
        "summary": summary,
        "records": results,
    }

    with open(
        RESULTS_PATH,
        "w",
        encoding="utf-8",
    ) as file:

        json.dump(
            output,
            file,
            indent=2,
            ensure_ascii=False,
        )

    print()
    print(
        f"Results saved to: {RESULTS_PATH}"
    )

def save_results_excel(results, summary):
    """
    Save the complete synthetic dataset together with four clean
    OPA-evaluation columns in one Excel sheet.

    Sheets:
        1. OPA Evaluation
        2. Summary
        3. Metadata

    The four additional columns are:
        1. Outcome
        2. Triggered Rules
        3. Explanation
        4. Suggested Remediation
    """

    from datetime import datetime
    from zoneinfo import ZoneInfo

    RESULTS_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    # ---------------------------------------------------------
    # Creation timestamp
    # ---------------------------------------------------------

    created_at = datetime.now(
        ZoneInfo("Asia/Kolkata")
    )

    created_date = created_at.strftime(
        "%d-%m-%Y"
    )

    created_time = created_at.strftime(
        "%I:%M:%S %p"
    )

    created_datetime = created_at.strftime(
        "%d-%m-%Y %I:%M:%S %p"
    )

    # ---------------------------------------------------------
    # Remediation by dataset column
    # ---------------------------------------------------------

    remediation_by_column = {
        "customer_name": "Remove or mask the customer's full name.",
        "email": "Remove or mask the personal email address.",
        "phone": "Remove or mask the phone number.",
        "address": "Remove or generalize the postal/home address.",
        "dob": "Remove or generalize the date of birth.",
        "gender": "Remove or generalize gender information where required.",
        "passport_number": "Remove the passport number completely.",
        "ni_number": "Remove the National Insurance number completely.",
        "credit_card_number": "Remove or mask the payment card number.",
        "bank_account": "Remove or mask the bank account number.",
        "medical_condition": "Remove the medical information.",
        "ethnicity": "Remove the ethnicity information.",
        "religion": "Remove the religion or belief information.",
        "political_view": "Remove the political opinion information.",
        "employee_id": "Remove or mask the employee identifier.",
        "department": "Remove or generalize the department when it contributes to identification.",
        "job_role": "Remove or generalize the job role when it contributes to identification.",
        "customer_id": "Remove or mask the customer identifier.",
        "ip_address": "Remove or mask the IP address.",
    }

    # ---------------------------------------------------------
    # Build one flat row per synthetic-data record
    # ---------------------------------------------------------

    rows = []

    for result in results:

        input_record = result.get(
            "input",
            {}
        )

        triggered_rules = result.get(
            "triggered_rules",
            []
        ) or []

        rule_ids = []
        explanations = []
        remediation_items = []
        seen_remediations = set()

        for rule in triggered_rules:

            rule_id = str(
                rule.get(
                    "rule_id",
                    ""
                )
            ).strip()

            if rule_id:
                rule_ids.append(rule_id)

            explanation = str(
                rule.get(
                    "explanation",
                    ""
                )
            ).strip()

            if explanation:
                explanations.append(
                    f"{rule_id}: {explanation}"
                    if rule_id
                    else explanation
                )

            columns = rule.get(
                "columns",
                []
            )

            if not isinstance(columns, list):
                columns = [columns]

            for column in columns:

                column = str(column).strip()

                if (
                    not column
                    or column not in input_record
                    or column not in remediation_by_column
                    or str(
                        input_record.get(
                            column,
                            ""
                        )
                    ).strip() == ""
                ):
                    continue

                if column in seen_remediations:
                    continue

                seen_remediations.add(column)

                remediation_items.append(
                    remediation_by_column[column]
                )

        row = {}

        # Entire original synthetic dataset
        for column in EXPECTED_COLUMNS:
            row[column] = input_record.get(
                column,
                ""
            )

        # -----------------------------------------------------
        # Exactly four additional evaluation columns
        # -----------------------------------------------------

        row["Outcome"] = result.get(
            "outcome",
            ""
        )

        row["Triggered Rules"] = (
            ", ".join(rule_ids)
            if rule_ids
            else "No violations"
        )

        row["Explanation"] = (
            " | ".join(explanations)
            if explanations
            else "No PII or sensitive information detected."
        )

        row["Suggested Remediation"] = (
            " | ".join(remediation_items)
            if remediation_items
            else "No remediation required."
        )

        rows.append(row)

    result_columns = EXPECTED_COLUMNS + [
        "Outcome",
        "Triggered Rules",
        "Explanation",
        "Suggested Remediation",
    ]

    results_dataframe = pd.DataFrame(
        rows,
        columns=result_columns,
    )

    # ---------------------------------------------------------
    # SUMMARY SHEET
    # ---------------------------------------------------------

    summary_dataframe = pd.DataFrame(
        [
            ["Total Records", summary["total_records"]],
            ["PASS", summary["pass"]],
            ["FLAG", summary["flag"]],
            ["BLOCK", summary["block"]],
            ["Pass Rate", f'{summary["pass_rate"]}%'],
        ],
        columns=["Metric", "Value"],
    )

    # ---------------------------------------------------------
    # METADATA SHEET
    # ---------------------------------------------------------

    metadata_dataframe = pd.DataFrame({
        "Property": [
            "Created Date",
            "Created Time",
            "Created Date & Time",
            "Timezone",
        ],
        "Value": [
            created_date,
            created_time,
            created_datetime,
            "Asia/Kolkata (IST)",
        ],
    })

    # ---------------------------------------------------------
    # WRITE WORKBOOK
    # ---------------------------------------------------------

    with pd.ExcelWriter(
        RESULTS_EXCEL_PATH,
        engine="openpyxl",
    ) as writer:

        results_dataframe.to_excel(
            writer,
            sheet_name="OPA Evaluation",
            index=False,
        )

        summary_dataframe.to_excel(
            writer,
            sheet_name="Summary",
            index=False,
        )

        metadata_dataframe.to_excel(
            writer,
            sheet_name="Metadata",
            index=False,
        )

    # ---------------------------------------------------------
    # FORMAT WORKBOOK
    # ---------------------------------------------------------

    workbook = load_workbook(
        RESULTS_EXCEL_PATH
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F2937",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    header_alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

    body_alignment = Alignment(
        vertical="top",
        wrap_text=True,
    )

    thin_border = Border(
        bottom=Side(
            style="thin",
            color="D1D5DB",
        )
    )

    outcome_fills = {
        "PASS": PatternFill(
            fill_type="solid",
            fgColor="DCFCE7",
        ),
        "FLAG": PatternFill(
            fill_type="solid",
            fgColor="FEF3C7",
        ),
        "BLOCK": PatternFill(
            fill_type="solid",
            fgColor="FEE2E2",
        ),
    }

    # ---------------------------------------------------------
    # OPA EVALUATION SHEET
    # ---------------------------------------------------------

    worksheet = workbook["OPA Evaluation"]

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    # Default readable widths for all columns
    for column_index, column in enumerate(
        result_columns,
        start=1,
    ):

        letter = get_column_letter(
            column_index
        )

        if column in {
            "Outcome",
        }:
            width = 14

        elif column in {
            "Triggered Rules",
        }:
            width = 28

        elif column in {
            "Explanation",
            "Suggested Remediation",
        }:
            width = 55

        elif column == "record_id":
            width = 12

        elif column == "feedback":
            width = 45

        else:
            width = 22

        worksheet.column_dimensions[
            letter
        ].width = width

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    for row in worksheet.iter_rows(
        min_row=2
    ):

        for cell in row:
            cell.alignment = body_alignment
            cell.border = thin_border

        outcome_cell = row[
            len(EXPECTED_COLUMNS)
        ]

        outcome_fill = outcome_fills.get(
            str(
                outcome_cell.value
            ).upper()
        )

        if outcome_fill:
            outcome_cell.fill = outcome_fill

    worksheet.row_dimensions[1].height = 34

    # ---------------------------------------------------------
    # SUMMARY SHEET
    # ---------------------------------------------------------

    worksheet = workbook["Summary"]

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    worksheet.column_dimensions[
        "A"
    ].width = 24

    worksheet.column_dimensions[
        "B"
    ].width = 22

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    for row in worksheet.iter_rows(
        min_row=2
    ):

        for cell in row:
            cell.alignment = body_alignment
            cell.border = thin_border

    # ---------------------------------------------------------
    # METADATA SHEET
    # ---------------------------------------------------------

    worksheet = workbook["Metadata"]

    worksheet.freeze_panes = "A2"

    worksheet.column_dimensions[
        "A"
    ].width = 25

    worksheet.column_dimensions[
        "B"
    ].width = 35

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    for row in worksheet.iter_rows(
        min_row=2
    ):

        for cell in row:
            cell.alignment = body_alignment
            cell.border = thin_border

    worksheet.row_dimensions[1].height = 25

    # ---------------------------------------------------------
    # SAVE WORKBOOK
    # ---------------------------------------------------------

    workbook.save(
        RESULTS_EXCEL_PATH
    )

    print(
        f"Clean OPA evaluation Excel saved to: "
        f"{RESULTS_EXCEL_PATH}"
    )

if __name__ == "__main__":
    main()